"""Builds the NTLS Networked Weather Station deck: photographs, with the timed script in the notes.

Glen, 2026-09-16: "You didn't provide timing marks or a script. You also didn't provide a powerpoint
with photographs."

Ten minutes, hard stop. John Wanda speaks; Gerald Knezek has the sending unit (Yagi, transceiver, LoRa,
sensors); Roger Wagner has the base receiver on his MakerPort. The clock is built backward from the
live link and leaves 6:30-7:30 unassigned as the retry.

EVERY SLIDE CARRIES ITS TIME MARK ON THE SLIDE and its spoken words in the speaker notes, so John can
read from presenter view and see where he should be. The photographs are the real ones from Glen's
"Sangala Initiative" album, downloaded to _Drafts\\Working\\NTLS Photos and chosen by looking at all 37:

    photo_24  the mountain wall, cultivated, with homesteads          the place
    photo_05  mud-and-wattle house with a thatched roof               how people live
    photo_06  a lorry crowded with standing passengers                how people travel
    photo_03  five boys on a red-dirt path, one with a panga          the farming family
    photo_18  a child crossing a brown torrent on a log               when the rain comes
    photo_34  a road dissolved into mud, people walking through it    what the rain does
    photo_37  students in red uniforms at a 3D printer and laptop     who will build it

python-pptx installed with `python -m pip install python-pptx` (no admin), the same way pymupdf was.
"""
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PHOTOS = (r"C:\Users\glenb\UVa Lab School Dropbox\AI Sandbox\Design through Making"
          r"\_Drafts\Working\NTLS Photos")
OUT = (r"C:\Users\glenb\UVa Lab School Dropbox\AI Sandbox\Design through Making"
       r"\_Drafts\Networked Weather Station (Ver 2.0).pptx")

W, H = Inches(13.333), Inches(7.5)
INK = RGBColor(0x1A, 0x1A, 0x1A)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
QUIET = RGBColor(0x8A, 0x8A, 0x8A)

# (photo | None, time mark, on-slide line, speaker notes)
SLIDES = [
    (None, "0:00",
     "The Networked Weather Station",
     "Mount Elgon, eastern Uganda.\n\nJOHN: \u201cThis is the Mount Elgon region of eastern Uganda, on "
     "the border with Kenya. In a moment you are going to watch a weather reading leave it.\u201d"),

    ("photo_24.jpg", "0:20", "Mount Elgon, eastern Uganda",
     "JOHN: \u201cPeople farm the slopes of an extinct volcano, on ground steep enough that you plant "
     "standing sideways. The soil is deep and rich. That is why they are there.\u201d"),

    ("photo_05.jpg", "0:45", "No grid. No internet.",
     "JOHN: \u201cThere is no electrical grid here. There is no internet. A house is built from the "
     "mountain itself \u2014 earth, poles and thatch.\u201d"),

    ("photo_06.jpg", "1:00", "",
     "JOHN: \u201cWhen people travel, they travel standing on the back of a lorry, on roads the rain "
     "takes out.\u201d\n\n(If the clock is ahead, this slide can be skipped without losing the case.)"),

    ("photo_03.jpg", "1:15", "One decision a year: when to plant",
     "JOHN: \u201cA family here makes one decision each year on which everything depends \u2014 when to "
     "plant. That used to be made from experience, because the rains came when they had always come. "
     "They no longer do. Planting a few weeks early or late now costs a family its year.\u201d"),

    ("photo_18.jpg", "1:40", "When enough rain falls, the slope moves",
     "JOHN: \u201cThe second problem is faster. When enough rain falls on a slope already saturated, "
     "the slope moves.\u201d"),

    ("photo_34.jpg", "1:55", "Mudslides kill people here every year \u2014 without warning",
     "JOHN: \u201cMudslides kill people on this mountain every year, and they arrive without warning. "
     "The reason there is no warning is simple. Nobody is measuring the rain.\u201d\n\nThen: \u201cSo "
     "that is what we are about to show you. Not a forecast made somewhere else \u2014 a measurement, "
     "taken where the people are, and carried out of a place with no infrastructure at all. "
     "Gerald.\u201d"),

    ("weather_station.png", "2:30", "The sending unit \u2014 Gerald",
     "GERALD (2:30\u20134:30): the weather sensors, the transceiver, and the Yagi antenna. LoRa radio "
     "\u2014 long range, very little power, and nothing in between. No tower, no subscription, no "
     "wire. Takes a reading."),

    ("weather_station.png", "4:30", "The base station \u2014 Roger",
     "ROGER (4:30\u20136:30): the receiver built on the MakerPort microcontroller. The reading Gerald "
     "just took arrives on screen.\n\n6:30\u20137:30 IS UNASSIGNED. If the link did not take, this is "
     "the retry. Use it and say nothing about it."),

    (None, "7:30", "One station is weather. A hundred is warning.",
     "JOHN: \u201cNow imagine that not once, but at a hundred points across a mountainside, each "
     "station passing what it measures to the next. One station tells you the weather. A mesh of them "
     "tells you where the rain is falling hardest and how long it has been falling there. That is the "
     "thing nobody in this region has today.\u201d"),

    ("photo_37.jpg", "8:00", "And these are the engineers who will build it",
     "JOHN: \u201cThe goal is not to go and install weather stations for them. It is a high school "
     "engineering program, so that students fabricate the stations, install them, and maintain them.\n\n"
     "A station that arrives from outside lasts until the day it breaks. One built by the person who "
     "lives beside it lasts as long as it is needed.\u201d\n\n(Last thing to cut. 8:40 \u2014 "
     "questions, and stop on time.)"),
]


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def textbox(slide, left, top, width, height, text, size, bold=False,
            color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Segoe UI"
    return box


def cover_image(slide, path):
    """Fill the slide with the photograph, cropping the overflow rather than distorting it."""
    from PIL import Image
    iw, ih = Image.open(path).size
    sw, sh = W / 914400.0, H / 914400.0                     # slide size in inches
    scale = max(sw / iw, sh / ih)
    w, h = iw * scale, ih * scale
    slide.shapes.add_picture(path, Inches((sw - w) / 2), Inches((sh - h) / 2),
                             Inches(w), Inches(h))


def fit_image(slide, path, bottom_clear):
    """CONTAIN rather than cover. The weather station is a wide flat-lay on white and the sensors sit
    at the far left and right - cropping it to 16:9 would cut them off, which is the whole subject."""
    from PIL import Image
    iw, ih = Image.open(path).size
    sw = W / 914400.0
    sh = H / 914400.0 - bottom_clear
    scale = min(sw / iw, sh / ih)
    w, h = iw * scale, ih * scale
    slide.shapes.add_picture(path, Inches((sw - w) / 2), Inches((sh - h) / 2 + 0.2),
                             Inches(w), Inches(h))


def scrim(slide, top, height):
    """A dark band behind the caption so white text reads over any photograph."""
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, top, W, height)
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor(0, 0, 0)
    s.fill.transparency = 0.35
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    for photo, mark, line, notes in SLIDES:
        slide = prs.slides.add_slide(blank)
        if photo:
            p = os.path.join(PHOTOS, photo)
            if not os.path.exists(p):
                sys.exit("missing photograph: " + p)
            flat = photo == "weather_station.png"
            if flat:
                fit_image(slide, p, 1.5)
                textbox(slide, Inches(0.6), H - Inches(1.25), W - Inches(1.2), Inches(0.9),
                        line, 30, True, INK)
                textbox(slide, W - Inches(1.5), Inches(0.25), Inches(1.1), Inches(0.4),
                        mark, 16, True, QUIET, PP_ALIGN.RIGHT)
            else:
                cover_image(slide, p)
                if line:
                    scrim(slide, H - Inches(1.35), Inches(1.35))
                    textbox(slide, Inches(0.6), H - Inches(1.15), W - Inches(1.2), Inches(0.9),
                            line, 30, True, PAPER)
                textbox(slide, W - Inches(1.5), Inches(0.25), Inches(1.1), Inches(0.4),
                        mark, 16, True, PAPER, PP_ALIGN.RIGHT)
        else:
            textbox(slide, Inches(1.0), Inches(2.7), W - Inches(2.0), Inches(2.0),
                    line, 44, True, INK)
            textbox(slide, W - Inches(1.5), Inches(0.25), Inches(1.1), Inches(0.4),
                    mark, 16, True, QUIET, PP_ALIGN.RIGHT)
        add_notes(slide, notes)

    if os.path.exists(OUT):
        sys.exit("refusing to overwrite " + OUT)
    prs.save(OUT)
    print("wrote", OUT)
    print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))


if __name__ == "__main__":
    main()
